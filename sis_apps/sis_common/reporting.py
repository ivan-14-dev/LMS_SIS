"""Safe allow-listed multi-format exports for tenant-configured reports."""

import csv
from io import BytesIO, StringIO

from django.http import HttpResponse
from django.utils import timezone
from django.utils.html import escape
from rest_framework.exceptions import ValidationError

from .academic_configuration import REPORT_EXPORT_FORMATS


def _has_any_permission(user, permissions):
    if not permissions:
        return True
    if getattr(user, "is_superuser", False):
        return True
    return any(user.has_perm(permission) for permission in permissions)


def configured_report(request, code, allowed_datasets=None):
    reports = getattr(request.tenant, "configuration_academique", {}).get("reports", [])
    report = next((item for item in reports if item.get("code") == code), None)
    if not report:
        raise ValidationError({"report": "Rapport inconnu ou non activé."})
    dataset = report.get("dataset")
    if allowed_datasets and dataset not in allowed_datasets:
        raise ValidationError({"report": "Ce rapport n'est pas disponible depuis ce module."})
    required_permissions = report.get("required_permissions", [])
    if required_permissions and not _has_any_permission(request.user, required_permissions):
        raise ValidationError({"report": "Vous n'avez pas accès à ce rapport."})
    return report


def _selected_format(report, supplied_format):
    export_format = str(supplied_format or "csv").strip().lower()
    allowed_formats = report.get("formats", ["csv"])
    if export_format not in REPORT_EXPORT_FORMATS:
        raise ValidationError(
            {"format": f"Format invalide. Valeurs acceptées: {', '.join(sorted(REPORT_EXPORT_FORMATS))}."}
        )
    if export_format not in allowed_formats:
        raise ValidationError({"format": "Ce format n'est pas autorisé pour ce rapport."})
    return export_format


def _filtered_rows(queryset, report, field_map, filter_map, supplied_filters):
    fields = report.get("fields", [])
    unknown_fields = set(fields) - set(field_map)
    if unknown_fields:
        raise ValidationError({"fields": f"Champs non exportables: {', '.join(sorted(unknown_fields))}."})
    allowed_filters = set(report.get("allowed_filters", []))
    unknown_filters = set(supplied_filters) - allowed_filters
    if unknown_filters:
        raise ValidationError({"filters": f"Filtres non autorisés: {', '.join(sorted(unknown_filters))}."})
    orm_filters = {}
    for name, value in supplied_filters.items():
        if name not in filter_map:
            raise ValidationError({"filters": f"Filtre inconnu: {name}."})
        orm_filters[filter_map[name]] = value
    queryset = queryset.filter(**orm_filters)
    headers = [field_map[field][0] for field in fields]
    orm_fields = [field_map[field][1] for field in fields]
    rows = list(queryset.values_list(*orm_fields).distinct())
    return headers, rows


def _report_metadata(request, report, supplied_filters):
    tenant = getattr(request, "tenant", None)
    filter_pairs = []
    included_filter_keys = set()
    for key in ("annee", "periode", "session", "semestre", "classe", "formation", "matiere", "ecue", "ue"):
        value = supplied_filters.get(key)
        if value not in (None, ""):
            included_filter_keys.add(key)
            filter_pairs.append((key.replace("_", " ").title(), value))
    extra_filters = {key: value for key, value in supplied_filters.items() if key not in included_filter_keys}
    metadata = [
        ("Établissement", getattr(tenant, "nom", "")),
        ("Type", tenant.get_type_display() if tenant and hasattr(tenant, "get_type_display") else getattr(tenant, "type", "")),
        ("Rapport", report.get("label", report.get("code", ""))),
        ("Dataset", report.get("dataset", "")),
        ("Généré le", timezone.now().strftime("%Y-%m-%d %H:%M:%S")),
        ("Couleur primaire", getattr(tenant, "couleur_primaire", "")),
        ("Couleur secondaire", getattr(tenant, "couleur_secondaire", "")),
    ]
    metadata.extend(filter_pairs)
    if extra_filters:
        metadata.append(
            (
                "Filtres",
                ", ".join(f"{key}={value}" for key, value in sorted(extra_filters.items())),
            )
        )
    return [(label, value) for label, value in metadata if value not in ("", None)]


def _csv_response(report, metadata, headers, rows):
    output = StringIO()
    writer = csv.writer(output)
    for label, value in metadata:
        writer.writerow([label, value])
    if metadata:
        writer.writerow([])
    writer.writerow(headers)
    writer.writerows(rows)
    response = HttpResponse(output.getvalue(), content_type="text/csv; charset=utf-8")
    response["Content-Disposition"] = f'attachment; filename="{report["code"]}.csv"'
    return response


def _xlsx_response(report, metadata, headers, rows):
    try:
        from openpyxl import Workbook
    except ImportError as exc:
        raise ValidationError({"format": "Le support Excel (.xlsx) n'est pas disponible."}) from exc

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = report.get("code", "report")[:31] or "report"
    current_row = 1
    for label, value in metadata:
        sheet.cell(row=current_row, column=1, value=label)
        sheet.cell(row=current_row, column=2, value=value)
        current_row += 1
    if metadata:
        current_row += 1
    for index, header in enumerate(headers, start=1):
        sheet.cell(row=current_row, column=index, value=header)
    current_row += 1
    for row in rows:
        for index, value in enumerate(row, start=1):
            sheet.cell(row=current_row, column=index, value=value)
        current_row += 1
    stream = BytesIO()
    workbook.save(stream)
    response = HttpResponse(
        stream.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    response["Content-Disposition"] = f'attachment; filename="{report["code"]}.xlsx"'
    return response


def _pdf_response(report, metadata, headers, rows):
    try:
        from weasyprint import HTML
    except ImportError as exc:
        raise ValidationError({"format": "Le support PDF n'est pas disponible."}) from exc

    metadata_rows = "".join(
        f"<tr><th>{escape(label)}</th><td>{escape(value)}</td></tr>" for label, value in metadata
    )
    table_rows_parts = []
    for row in rows:
        row_cells = "".join(f"<td>{escape(value)}</td>" for value in row)
        table_rows_parts.append(f"<tr>{row_cells}</tr>")
    table_rows = "".join(table_rows_parts)
    html = f"""
    <html>
      <head>
        <meta charset="utf-8" />
        <style>
          body {{ font-family: sans-serif; color: #1f2933; font-size: 12px; }}
          h1 {{ margin-bottom: 12px; }}
          table {{ width: 100%; border-collapse: collapse; margin-bottom: 16px; }}
          th, td {{ border: 1px solid #d9e2ec; padding: 6px 8px; text-align: left; vertical-align: top; }}
          th {{ background: #f5f7fa; }}
        </style>
      </head>
      <body>
        <h1>{escape(report.get("label", report.get("code", "Rapport")))}</h1>
        <table>{metadata_rows}</table>
        <table>
          <thead>
            <tr>{"".join(f"<th>{escape(header)}</th>" for header in headers)}</tr>
          </thead>
          <tbody>{table_rows}</tbody>
        </table>
      </body>
    </html>
    """
    response = HttpResponse(content_type="application/pdf")
    response["Content-Disposition"] = f'attachment; filename="{report["code"]}.pdf"'
    HTML(string=html).write_pdf(response)
    return response


def export_queryset(request, queryset, report, field_map, filter_map, supplied_filters):
    export_format = _selected_format(report, request.data.get("format"))
    headers, rows = _filtered_rows(queryset, report, field_map, filter_map, supplied_filters)
    metadata = _report_metadata(request, report, supplied_filters)
    if export_format == "xlsx":
        return _xlsx_response(report, metadata, headers, rows)
    if export_format == "pdf":
        return _pdf_response(report, metadata, headers, rows)
    return _csv_response(report, metadata, headers, rows)
