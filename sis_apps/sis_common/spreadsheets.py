"""Helpers for strict Excel import/export flows."""

from io import BytesIO

from django.http import HttpResponse

from rest_framework.exceptions import ValidationError

FORMULA_PREFIXES = ("=", "+", "-", "@")


def _normalized_headers(headers):
    return [str(header or "").strip() for header in headers]


def _reject_formula_like_value(value, field_name):
    if isinstance(value, str) and value.lstrip().startswith(FORMULA_PREFIXES):
        raise ValidationError(
            {field_name: "Les formules et expressions dynamiques ne sont pas autorisées."}
        )


def _load_xlsx_rows(uploaded_file):
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise ValidationError({"file": "Le support des imports Excel (.xlsx) n'est pas disponible."}) from exc
    workbook = load_workbook(filename=BytesIO(uploaded_file.read()), read_only=True, data_only=True)
    worksheet = workbook.active
    return [list(row) for row in worksheet.iter_rows(values_only=True)]


def _load_xls_rows(uploaded_file):
    try:
        import xlrd
    except ImportError as exc:
        raise ValidationError({"file": "Le support des imports Excel (.xls) n'est pas disponible."}) from exc
    workbook = xlrd.open_workbook(file_contents=uploaded_file.read())
    sheet = workbook.sheet_by_index(0)
    return [sheet.row_values(index) for index in range(sheet.nrows)]


def load_excel_rows(uploaded_file, expected_headers):
    filename = getattr(uploaded_file, "name", "").lower()
    if filename.endswith(".xlsx"):
        rows = _load_xlsx_rows(uploaded_file)
    elif filename.endswith(".xls"):
        rows = _load_xls_rows(uploaded_file)
    else:
        raise ValidationError({"file": "Seuls les fichiers .xlsx et .xls sont autorisés."})
    if not rows:
        raise ValidationError({"file": "Le fichier importé est vide."})
    headers = _normalized_headers(rows[0])
    expected = _normalized_headers(expected_headers)
    if headers != expected:
        raise ValidationError(
            {
                "columns": (
                    "Les colonnes attendues sont: {}.".format(", ".join(expected))
                )
            }
        )
    normalized_rows = []
    for row_number, raw_row in enumerate(rows[1:], start=2):
        values = list(raw_row[: len(expected)]) + [None] * max(0, len(expected) - len(raw_row))
        if all(value in (None, "") for value in values):
            continue
        row = {}
        for index, header in enumerate(expected):
            value = values[index]
            if isinstance(value, str):
                value = value.strip()
                _reject_formula_like_value(value, header)
            row[header] = value
        row["__row_number__"] = row_number
        normalized_rows.append(row)
    return normalized_rows


def template_response(headers, filename, sample_row=None):
    try:
        from openpyxl import Workbook
    except ImportError as exc:
        raise ValidationError({"format": "Le support Excel (.xlsx) n'est pas disponible."}) from exc
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.append(headers)
    if sample_row:
        worksheet.append(sample_row)
    stream = BytesIO()
    workbook.save(stream)
    response = HttpResponse(
        stream.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    response["Content-Disposition"] = f'attachment; filename="{filename}"'
    return response
